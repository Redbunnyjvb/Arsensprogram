"""Excel workbook <-> Project (openpyxl). Sheets: Project, Models, Tags, Sensors, Validation.

JSON is the source of truth; Excel is the human-friendly authoring surface (no opaque AR blobs).
"""
from __future__ import annotations

from openpyxl import Workbook, load_workbook

from .models import (
    CoordinateFrame,
    Marker,
    OriginCorner,
    PlacementOrigin,
    Project,
    Sensor,
    SensorStatus,
    StlModel,
)

MODEL_COLS = ["id", "name", "file_name", "role", "scale_percent",
              "offset_x", "offset_y", "offset_z", "rot_x", "rot_y", "rot_z", "visible"]
TAG_COLS = ["id", "type", "size_mm", "x_mm", "y_mm", "z_mm",
            "rot_x", "rot_y", "rot_z", "active", "pose_weight", "origin"]
SENSOR_COLS = ["order", "id", "name", "side", "x_mm", "y_mm", "z_mm", "nx", "ny", "nz",
               "tolerance_mm", "instruction", "status", "origin", "reference_tag_id", "sensor_tag_id"]


def _bool(v, default=False):
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in ("1", "true", "ja", "yes", "waar")


def _int(v, default=0):
    if v is None or str(v).strip() == "":
        return default
    return int(float(v))


def _float(v, default=0.0):
    if v is None or str(v).strip() == "":
        return default
    return float(v)


def _opt_int(v):
    if v is None or str(v).strip() == "":
        return None
    return int(float(v))


def write_workbook(project: Project, path, issues=None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Project"
    ws.append(["key", "value"])
    cf = project.coordinate_frame
    for key, value in [
        ("project_name", project.project_name),
        ("model_file", project.model_file),
        ("dim_x_mm", project.dimensions_mm[0]),
        ("dim_y_mm", project.dimensions_mm[1]),
        ("dim_z_mm", project.dimensions_mm[2]),
        ("dimensions_locked", project.dimensions_locked),
        ("origin_corner", cf.origin_corner.value),
        ("flip_x", cf.flip_x),
        ("flip_y", cf.flip_y),
        ("flip_z", cf.flip_z),
    ]:
        ws.append([key, value])

    wm = wb.create_sheet("Models")
    wm.append(MODEL_COLS)
    for m in project.stl_models:
        wm.append([m.id, m.name, m.file_name, m.role, m.scale_percent,
                   m.offset_mm[0], m.offset_mm[1], m.offset_mm[2],
                   m.rotation_deg[0], m.rotation_deg[1], m.rotation_deg[2], m.visible])

    wt = wb.create_sheet("Tags")
    wt.append(TAG_COLS)
    for t in project.markers:
        wt.append([t.id, t.type, t.size_mm, t.position_mm[0], t.position_mm[1], t.position_mm[2],
                   t.rotation_deg[0], t.rotation_deg[1], t.rotation_deg[2], t.active, t.pose_weight, t.origin.value])

    wsn = wb.create_sheet("Sensors")
    wsn.append(SENSOR_COLS)
    for s in project.sensors:
        wsn.append([s.order, s.id, s.name, s.side, s.position_mm[0], s.position_mm[1], s.position_mm[2],
                    s.normal[0], s.normal[1], s.normal[2], s.tolerance_mm, s.instruction,
                    s.status.value, s.origin.value,
                    "" if s.reference_tag_id is None else s.reference_tag_id,
                    "" if s.sensor_tag_id is None else s.sensor_tag_id])

    wv = wb.create_sheet("Validation")
    wv.append(["level", "message"])
    for issue in (issues or []):
        wv.append([getattr(issue, "level", ""), getattr(issue, "message", str(issue))])

    wb.save(path)


def _header_index(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, []
    head = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    return {h: i for i, h in enumerate(head)}, rows[1:]


def _cell(row, idx, key, default=None):
    i = idx.get(key)
    if i is None or i >= len(row):
        return default
    v = row[i]
    return default if v is None else v


def read_workbook(path) -> Project:
    wb = load_workbook(path, data_only=True)

    pdata = {}
    if "Project" in wb.sheetnames:
        for row in wb["Project"].iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                pdata[str(row[0])] = row[1]
    frame = CoordinateFrame(
        origin_corner=OriginCorner(str(pdata.get("origin_corner", "front_left_bottom"))),
        flip_x=_bool(pdata.get("flip_x")), flip_y=_bool(pdata.get("flip_y")), flip_z=_bool(pdata.get("flip_z")),
    )

    models = []
    if "Models" in wb.sheetnames:
        idx, body = _header_index(wb["Models"])
        for r in body:
            if r is None or all(c is None for c in r):
                continue
            models.append(StlModel(
                id=str(_cell(r, idx, "id", "")), name=str(_cell(r, idx, "name", "") or ""),
                file_name=str(_cell(r, idx, "file_name", "") or ""),
                role=str(_cell(r, idx, "role", "OTHER") or "OTHER"),
                scale_percent=_int(_cell(r, idx, "scale_percent", 100), 100),
                offset_mm=[_int(_cell(r, idx, "offset_x")), _int(_cell(r, idx, "offset_y")), _int(_cell(r, idx, "offset_z"))],
                rotation_deg=[_int(_cell(r, idx, "rot_x")), _int(_cell(r, idx, "rot_y")), _int(_cell(r, idx, "rot_z"))],
                visible=_bool(_cell(r, idx, "visible", True), True)))

    tags = []
    if "Tags" in wb.sheetnames:
        idx, body = _header_index(wb["Tags"])
        for r in body:
            if r is None or all(c is None for c in r):
                continue
            tags.append(Marker(
                id=_int(_cell(r, idx, "id", 0)), type=str(_cell(r, idx, "type", "apriltag") or "apriltag"),
                size_mm=_int(_cell(r, idx, "size_mm", 100), 100),
                position_mm=[_int(_cell(r, idx, "x_mm")), _int(_cell(r, idx, "y_mm")), _int(_cell(r, idx, "z_mm"))],
                rotation_deg=[_float(_cell(r, idx, "rot_x")), _float(_cell(r, idx, "rot_y")), _float(_cell(r, idx, "rot_z"))],
                active=_bool(_cell(r, idx, "active", True), True),
                pose_weight=_float(_cell(r, idx, "pose_weight", 1.0), 1.0),
                origin=PlacementOrigin(str(_cell(r, idx, "origin", "prepared") or "prepared"))))

    sensors = []
    if "Sensors" in wb.sheetnames:
        idx, body = _header_index(wb["Sensors"])
        for r in body:
            if r is None or all(c is None for c in r):
                continue
            sensors.append(Sensor(
                order=_int(_cell(r, idx, "order", 0)), id=str(_cell(r, idx, "id", "")),
                name=str(_cell(r, idx, "name", "") or ""), side=str(_cell(r, idx, "side", "veld") or "veld"),
                position_mm=[_int(_cell(r, idx, "x_mm")), _int(_cell(r, idx, "y_mm")), _int(_cell(r, idx, "z_mm"))],
                normal=[_float(_cell(r, idx, "nx", 0)), _float(_cell(r, idx, "ny", 1), 1.0), _float(_cell(r, idx, "nz", 0))],
                tolerance_mm=_int(_cell(r, idx, "tolerance_mm", 50), 50),
                instruction=str(_cell(r, idx, "instruction", "") or ""),
                status=SensorStatus(str(_cell(r, idx, "status", "pending") or "pending")),
                origin=PlacementOrigin(str(_cell(r, idx, "origin", "prepared") or "prepared")),
                reference_tag_id=_opt_int(_cell(r, idx, "reference_tag_id")),
                sensor_tag_id=_opt_int(_cell(r, idx, "sensor_tag_id"))))

    return Project(
        project_name=str(pdata.get("project_name", "Transformer A")),
        model_file=str(pdata.get("model_file", "transformer_model.glb")),
        dimensions_mm=[_int(pdata.get("dim_x_mm", 10000), 10000),
                       _int(pdata.get("dim_y_mm", 5000), 5000),
                       _int(pdata.get("dim_z_mm", 3200), 3200)],
        dimensions_locked=_bool(pdata.get("dimensions_locked", False)),
        coordinate_frame=frame,
        sensors=sorted(sensors, key=lambda s: s.order),
        markers=tags,
        stl_models=models,
    )
