"""Adapter: Excel authoring workbook <-> model (openpyxl, lazily imported).

Authoring layer: carries the human-editable fields only. Opaque AR blobs (placement, drift,
coordinate_frame) are NOT represented in Excel — use JSON/packages for lossless round-trips.
"""
from __future__ import annotations

from .model import (
    FloatVector,
    Marker,
    MmPosition,
    PlacementOrigin,
    Project,
    Sensor,
    SensorStatus,
)

SENSOR_HEADERS = [
    "order", "id", "name", "side", "x_mm", "y_mm", "z_mm",
    "tolerance_mm", "instruction", "sensor_tag_id", "origin", "status",
]
TAG_HEADERS = [
    "id", "type", "size_mm", "x_mm", "y_mm", "z_mm",
    "rot_x", "rot_y", "rot_z", "origin", "active",
]


def _require_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("Excel-ondersteuning vereist openpyxl (pip install openpyxl).") from exc


def _origin(value) -> PlacementOrigin:
    for o in PlacementOrigin:
        if o.value == str(value):
            return o
    return PlacementOrigin.PREPARED


def _status(value) -> SensorStatus:
    for s in SensorStatus:
        if s.value == str(value):
            return s
    return SensorStatus.PENDING


def _bool(value, default=False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "ja", "yes", "waar")


def _opt_int(value):
    if value is None or str(value).strip() == "":
        return None
    return int(float(value))


def _int(value, default=0) -> int:
    if value is None or str(value).strip() == "":
        return default
    return int(float(value))


def _float(value, default=0.0) -> float:
    if value is None or str(value).strip() == "":
        return default
    return float(value)


def write_workbook(project: Project, path) -> None:
    openpyxl = _require_openpyxl()
    wb = openpyxl.Workbook()

    ws_p = wb.active
    ws_p.title = "Project"
    ws_p.append(["key", "value"])
    ws_p.append(["project_name", project.project_name])
    ws_p.append(["model_file", project.model_file])
    ws_p.append(["dim_x_mm", project.dimensions_mm.x])
    ws_p.append(["dim_y_mm", project.dimensions_mm.y])
    ws_p.append(["dim_z_mm", project.dimensions_mm.z])
    ws_p.append(["dimensions_locked", project.dimensions_locked])

    ws_s = wb.create_sheet("Sensors")
    ws_s.append(SENSOR_HEADERS)
    for s in project.sensors:
        ws_s.append([
            s.order, s.id, s.name, s.side,
            s.position_mm.x, s.position_mm.y, s.position_mm.z,
            s.tolerance_mm, s.instruction,
            "" if s.sensor_tag_id is None else s.sensor_tag_id,
            s.origin.value, s.status.value,
        ])

    ws_t = wb.create_sheet("Tags")
    ws_t.append(TAG_HEADERS)
    for m in project.markers:
        ws_t.append([
            m.id, m.type, m.size_mm,
            m.position_mm.x, m.position_mm.y, m.position_mm.z,
            m.rotation_deg.x, m.rotation_deg.y, m.rotation_deg.z,
            m.origin.value, m.active,
        ])

    wb.save(path)


def _header_index(rows) -> dict:
    if not rows:
        return {}
    head = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    return {h: i for i, h in enumerate(head)}


def _cell(row, idx: dict, key: str, default=None):
    i = idx.get(key)
    if i is None or i >= len(row):
        return default
    v = row[i]
    return default if v is None else v


def read_workbook(path) -> Project:
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=True)

    pdata: dict = {}
    if "Project" in wb.sheetnames:
        for row in wb["Project"].iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                pdata[str(row[0])] = row[1]

    sensors: list[Sensor] = []
    if "Sensors" in wb.sheetnames:
        rows = list(wb["Sensors"].iter_rows(values_only=True))
        idx = _header_index(rows)
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            sensors.append(Sensor(
                order=_int(_cell(r, idx, "order", 0)),
                id=str(_cell(r, idx, "id", "") or ""),
                name=str(_cell(r, idx, "name", "") or ""),
                side=str(_cell(r, idx, "side", "veld") or "veld"),
                position_mm=MmPosition(_int(_cell(r, idx, "x_mm")), _int(_cell(r, idx, "y_mm")), _int(_cell(r, idx, "z_mm"))),
                tolerance_mm=_int(_cell(r, idx, "tolerance_mm", 50), 50),
                instruction=str(_cell(r, idx, "instruction", "") or ""),
                sensor_tag_id=_opt_int(_cell(r, idx, "sensor_tag_id")),
                origin=_origin(_cell(r, idx, "origin", "prepared")),
                status=_status(_cell(r, idx, "status", "pending")),
            ))

    markers: list[Marker] = []
    if "Tags" in wb.sheetnames:
        rows = list(wb["Tags"].iter_rows(values_only=True))
        idx = _header_index(rows)
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            markers.append(Marker(
                id=_int(_cell(r, idx, "id", 0)),
                type=str(_cell(r, idx, "type", "apriltag") or "apriltag"),
                size_mm=_int(_cell(r, idx, "size_mm", 100), 100),
                position_mm=MmPosition(_int(_cell(r, idx, "x_mm")), _int(_cell(r, idx, "y_mm")), _int(_cell(r, idx, "z_mm"))),
                rotation_deg=FloatVector(_float(_cell(r, idx, "rot_x")), _float(_cell(r, idx, "rot_y")), _float(_cell(r, idx, "rot_z"))),
                origin=_origin(_cell(r, idx, "origin", "prepared")),
                active=_bool(_cell(r, idx, "active", True), True),
            ))

    return Project(
        project_name=str(pdata.get("project_name", "Transformer A")),
        model_file=str(pdata.get("model_file", "transformer_model.glb")),
        dimensions_mm=MmPosition(
            _int(pdata.get("dim_x_mm", 10000), 10000),
            _int(pdata.get("dim_y_mm", 5000), 5000),
            _int(pdata.get("dim_z_mm", 3200), 3200),
        ),
        dimensions_locked=_bool(pdata.get("dimensions_locked", False)),
        sensors=sorted(sensors, key=lambda s: s.order),
        markers=markers,
    )
